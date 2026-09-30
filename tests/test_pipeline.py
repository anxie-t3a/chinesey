"""Tests for the vocab pipeline. They run offline: Claude and Google Sheets are replaced by fakes.

Run with:  poetry run pytest
Also run by .github/workflows/publish_changes_to_source_of_truth.yml before anything is written to the sheet.
"""
import copy
import io
import json
import types

import openpyxl

import build
import review
import sync
import writeback
from common import SCHEMA, card_id
from diff import diff, find_duplicates
from edits import apply_edits, derive_edits
from sheet import read_tabs, validate
from vocab import build_lexicon, unknown_parts, uses_target


def test_card_id_separates_same_hanzi(row):
    """Same hanzi with a different meaning is a different card; notes don't affect the id."""
    mw = row(2, "家", "jiā", "measure word for businesses", "measure_word")
    home = row(3, "家", "jiā", "home")
    assert card_id(mw) != card_id(home)
    assert card_id(home) == card_id(row(9, "家", "jiā", "home", notes="different notes"))


def test_diff_classifies_rows(row):
    """Each kind of sheet change is classified correctly, and identity changes produce aliases."""
    prev = {"Food": [row(2, "米饭", "mǐfàn", "rice"), row(3, "面条", "miàntiáo", "noodles"),
                     row(4, "汤", "tāng", "soup"), row(5, "饺子", "jiǎozi", "dumplings")]}
    curr = {"Food": [row(2, "米饭", "mǐfàn", "cooked rice"),                        # core change, same hanzi
                     row(3, "面条", "miàntiáo", "noodles", notes="usually wheat"),  # non-core change
                     row(5, "饺子儿", "jiǎozi", "dumplings"),                       # hanzi edit, same row
                     row(6, "茶", "chá", "tea")]}                                   # new; 汤 removed
    d = diff(prev, curr)
    changed = {c["hanzi"]: c for c in d["changed"]}
    assert changed["米饭"]["changes"] == {"english": ["rice", "cooked rice"]}
    assert changed["米饭"]["old_id"] != changed["米饭"]["id"]
    assert changed["面条"]["id"] == changed["面条"]["old_id"]
    assert "饺子儿" in changed
    assert [n["hanzi"] for n in d["new"]] == ["茶"]
    assert [r["hanzi"] for r in d["removed"]] == ["汤"]
    assert d["aliases"] == {changed["米饭"]["old_id"]: changed["米饭"]["id"],
                            changed["饺子儿"]["old_id"]: changed["饺子儿"]["id"]}


def test_find_duplicates(row):
    """Exact duplicates, partial copies, conflicting duplicates and same-hanzi rows are told apart."""
    tabs = {"Verbs": [row(2, "跑", "pǎo", "run", "verb"), row(3, "跑", "pǎo", "run", "verb"),
                      row(4, "走", "zǒu", "walk", "verb", notes="a"), row(5, "走", "zǒu", "walk", "verb", notes="b"),
                      row(6, "健身", "jiànshēn", "work out", "VOV"), row(7, "健身", "jiànshēn", "", ""),
                      row(8, "行", "xíng", "OK", "adjective"), row(9, "行", "háng", "row", "measure_word")]}
    d = find_duplicates(tabs)
    assert [(x["hanzi"], x["dup_row"], x.get("partial", False)) for x in d["exact"]] == [
        ("跑", 3, False), ("健身", 7, True)]
    assert [(x["hanzi"], x["rows"]) for x in d["conflicts"]] == [("走", [4, 5])]
    assert [(x["hanzi"], x["rows"]) for x in d["similar"]] == [("行", [8, 9])]


def test_validate_blocks_bad_rows(row):
    """Rows with errors are blocked; warnings don't block."""
    tabs = {"Objects": [row(2, "纸", "zhǐ", "paper", measure_word="张 (zhāng)"),
                        row(3, "笔", "bǐ", "pen", type="thing"),
                        row(4, "书", "shu1", "book", measure_word="933")]}
    issues, blocked = validate(tabs)
    assert blocked == {("Objects", 3), ("Objects", 4)}
    assert {(i["row"], i["field"], i["level"]) for i in issues} == {
        (3, "type", "error"), (4, "measure_word", "error"), (4, "pinyin", "warning")}


def test_vocab_check():
    """Sentence checking: segmentation into known words, and finding the card's own word."""
    lex = build_lexicon(["咖啡", "喝", "要是……就……", "理发店 / 美发店"], {"我", "不", "了"})
    assert unknown_parts("我不喝咖啡。", lex) == []
    assert unknown_parts("我喝茶了", lex) == ["茶"]
    assert uses_target("要是下雨就不去", "要是……就……")
    assert not uses_target("就要是", "要是……就……")
    assert uses_target("我去美发店", "理发店 / 美发店")


def _tabs(row):
    """Two tabs sharing the card 地铁, plus 杯子 in one tab."""
    return {"Objects": [row(2, "地铁", "dìtiě", "subway"), row(3, "杯子", "bēizi", "cup")],
            "Places": [row(2, "地铁", "dìtiě", "subway", notes="take the subway")]}


def test_review_edits_hit_every_membership_and_respect_rejections(row):
    """Review fixes go to every row of a card, invalid suggestions are dropped, and rejections stick."""
    tabs = _tabs(row)
    subway, cup = card_id(tabs["Objects"][0]), card_id(tabs["Objects"][1])
    reviews = {
        subway: {"issues": [{"field": "measure_word", "kind": "missing", "current": "",
                             "suggested": "条 (tiáo)", "reason": "lines", "confidence": "high"}],
                 "example": {"hanzi": "我坐地铁", "pinyin": "wǒ zuò dìtiě", "english": "I take the subway"}},
        cup: {"issues": [{"field": "type", "kind": "incorrect", "current": "noun", "suggested": "thing",
                          "reason": "invalid suggestion is dropped", "confidence": "high"}], "example": None},
    }
    dups = find_duplicates(tabs)
    edits = derive_edits(tabs, set(), dups, reviews, {}, [])
    mw_edits = [e for e in edits if e["field"] == "measure_word"]
    assert sorted(e["tab"] for e in mw_edits) == ["Objects", "Places"]
    assert [e["after"] for e in edits if e["field"] == "examples"] == ["我坐地铁 | wǒ zuò dìtiě | I take the subway"]
    assert not [e for e in edits if e["field"] == "type"]

    carried = [{**mw_edits[0], "status": "rejected"}]
    again = derive_edits(tabs, set(), dups, reviews, {}, carried)
    assert {e["key"]: e["status"] for e in again}[mw_edits[0]["key"]] == "rejected"
    decided = derive_edits(tabs, set(), dups, reviews, {mw_edits[0]["key"]: {"status": "rejected"}}, [])
    assert mw_edits[0]["key"] not in {e["key"] for e in decided}

    edited, aliases = apply_edits(tabs, again)
    assert edited["Objects"][0]["measure_word"] == ""   # rejected edit not applied
    assert edited["Places"][0]["measure_word"] == "条 (tiáo)"
    assert aliases == {}


def test_identity_edit_creates_alias(row):
    """Editing the english gloss changes the id, and the old id becomes an alias."""
    tabs = _tabs(row)
    old = card_id(tabs["Objects"][1])
    reviews = {old: {"issues": [{"field": "english", "kind": "incorrect", "current": "cup",
                                 "suggested": "cup / glass", "reason": "", "confidence": "medium"}],
                     "example": None}}
    edits = derive_edits(tabs, set(), find_duplicates(tabs), reviews, {}, [])
    edited, aliases = apply_edits(tabs, edits)
    cards, _ = build.build_cards(edited, set(), aliases)
    cup = next(c for c in cards["cards"] if c["hanzi"] == "杯子")
    assert cup["english"] == "cup / glass" and cup["aliases"] == [old]


class FakeClient:
    """Stands in for anthropic.Anthropic; replies with queued JSON payloads."""

    def __init__(self, replies):
        # replies: JSON payloads returned in order, one per API call
        self.replies, self.requests = list(replies), []
        self.messages = self

    def stream(self, **kwargs):
        # Mimics client.messages.stream(...): records the request, returns a context manager.
        self.requests.append(kwargs)
        text = json.dumps(self.replies.pop(0), ensure_ascii=False)
        message = types.SimpleNamespace(stop_reason="end_turn",
                                        content=[types.SimpleNamespace(type="text", text=text)])
        return _Ctx(types.SimpleNamespace(get_final_message=lambda: message))


class _Ctx:
    """Minimal context manager wrapping a value (stands in for the SDK's stream object)."""
    def __init__(self, value):
        self.value = value

    def __enter__(self):
        return self.value

    def __exit__(self, *exc):
        return False


def test_reviewer_retries_rejected_example(row, monkeypatch):
    """An example with a word outside the app vocab is rejected and retried with feedback."""
    tabs = _tabs(row)
    cards, _ = build.build_cards(tabs, set(), {})
    cards = cards["cards"]
    bad = {"cards": [{"ref": 0, "issues": [], "example": {"hanzi": "我坐地铁上班", "pinyin": "x", "english": "x"}},
                     {"ref": 1, "issues": [], "example": {"hanzi": "我的杯子", "pinyin": "wǒ de bēizi",
                                                          "english": "my cup"}}]}
    good = {"cards": [{"ref": 0, "issues": [], "example": {"hanzi": "我不坐地铁", "pinyin": "wǒ bú zuò dìtiě",
                                                           "english": "I don't take the subway"}}]}
    client = FakeClient([bad, good])
    lexicon = build_lexicon([c["hanzi"] for c in cards], {"我", "的", "不", "坐"})
    reviewer = review.Reviewer(cards, {"我", "的", "不", "坐"}, lexicon, client=client)
    records, note = reviewer.review(cards)
    assert note == ""
    assert records[cards[0]["id"]]["example"]["hanzi"] == "我不坐地铁"
    assert records[cards[1]["id"]]["example"]["hanzi"] == "我的杯子"
    assert "上班" in json.dumps(client.requests[1]["messages"], ensure_ascii=False)
    assert client.requests[0]["model"] == "claude-opus-5-5"


def _xlsx(tabs) -> bytes:
    """Build a real .xlsx file from rows, including a `_readme` tab that must be ignored."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in tabs.items():
        ws = wb.create_sheet(name)
        ws.append(SCHEMA)
        for r in rows:
            while ws.max_row < r["row"] - 1:
                ws.append([None])
            ws.append([r[f] or None for f in SCHEMA])
    wb.create_sheet("_readme").append(["ignored"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_tabs_roundtrip(row):
    """Rows written to xlsx read back identically."""
    tabs = _tabs(row)
    parsed, errors = read_tabs(_xlsx(tabs))
    assert errors == [] and parsed == tabs


class FakeWorksheet:
    """Stands in for a gspread worksheet: holds cell values in memory."""
    def __init__(self, rows):
        self.values = [SCHEMA] + [[r[f] for f in SCHEMA] for r in rows]

    def get_all_values(self):
        return copy.deepcopy(self.values)

    def batch_update(self, updates, value_input_option):
        for u in updates:
            cell = u["range"].split(":")[0]
            col, n = ord(cell[0]) - 65, int(cell[1:])
            for i, v in enumerate(u["values"][0]):
                self.values[n - 1][col + i] = v


def test_sync_then_writeback(row, tmp_path, monkeypatch):
    """Full cycle: sync proposes removing a duplicate, write-back clears it in the sheet, next sync is clean."""
    # Point every data file at a temp dir so the test never touches the real data/.
    for mod in (build, sync, writeback):
        for name in ("SNAPSHOT_FILE", "EDITS_FILE", "ALIASES_FILE", "DECISIONS_FILE", "CARDS_FILE"):
            if hasattr(mod, name):
                monkeypatch.setattr(mod, name, tmp_path / getattr(build, name).name)
    monkeypatch.setattr(sync, "REVIEWS_FILE", tmp_path / "reviews.json")
    monkeypatch.setattr(sync, "REPORTS", tmp_path / "reports")

    tabs = _tabs(row)
    tabs["Objects"].append(row(4, "杯子", "bēizi", "cup"))  # exact duplicate
    xlsx = tmp_path / "v.xlsx"
    xlsx.write_bytes(_xlsx(tabs))
    assert sync.main(["--xlsx", str(xlsx), "--no-review"]) == 0
    edits = json.loads((tmp_path / "sheet_edits.json").read_text())
    assert [(e["op"], e["row"]) for e in edits] == [("clear_row", 4)]
    cards = json.loads((tmp_path / "cards.json").read_text())
    assert len(cards["cards"]) == 2

    sheet = {"Objects": FakeWorksheet(tabs["Objects"]), "Places": FakeWorksheet(tabs["Places"])}
    book = types.SimpleNamespace(worksheet=sheet.__getitem__)
    fake_gspread = types.SimpleNamespace(
        service_account_from_dict=lambda info: types.SimpleNamespace(open_by_key=lambda key: book))
    monkeypatch.setitem(__import__("sys").modules, "gspread", fake_gspread)
    monkeypatch.setenv("GOOGLE_SERVICE_ACCOUNT_JSON", "{}")
    assert writeback.main(["--sheet-id", "x"]) == 0
    assert sheet["Objects"].values[3] == [""] * len(SCHEMA)
    assert json.loads((tmp_path / "sheet_edits.json").read_text()) == []
    decisions = json.loads((tmp_path / "edit_decisions.json").read_text())
    assert [d["status"] for d in decisions.values()] == ["applied"]

    # Next sync sees the cleaned sheet: no diff noise, nothing left to propose.
    tabs["Objects"].pop()
    xlsx.write_bytes(_xlsx(tabs))
    assert sync.main(["--xlsx", str(xlsx), "--no-review"]) == 0
    report = (tmp_path / "reports" / "sync-report.md").read_text()
    assert "| removed rows | 0 |" in report and "| proposed sheet edits | 0 |" in report


def test_writeback_skips_cells_changed_since_pr(row):
    """Write-back refuses to overwrite a changed cell, but finds a row that merely moved."""
    ws = FakeWorksheet([row(2, "杯子", "bēizi", "mug")])
    edit = {"op": "set", "row": 2, "hanzi": "杯子", "field": "english", "before": "cup", "after": "cup / glass"}
    assert writeback.locate(ws.get_all_values(), edit) is None
    moved = FakeWorksheet([row(2, "茶", "chá", "tea"), row(3, "杯子", "bēizi", "cup")])
    assert writeback.locate(moved.get_all_values(), edit) == 3


def test_review_selection_only_changed_unless_review_all(row):
    """Only cards whose content changed since their review are re-reviewed; review_all forces every card."""
    tabs = _tabs(row)
    cards = build.build_cards(tabs, set(), {})[0]["cards"]
    reviews = {c["id"]: {"hash": review.review_hash(c)} for c in cards}
    assert review.needs_review(cards, reviews) == []
    assert len(review.needs_review(cards, reviews, review_all=True)) == len(cards)

    tabs["Objects"][1]["notes"] = "usually for drinks"   # a notes-only change counts as changed content
    cards = build.build_cards(tabs, set(), {})[0]["cards"]
    assert [c["hanzi"] for c in review.needs_review(cards, reviews)] == ["杯子"]


def test_sync_reviews_changed_cards_and_review_all_flag(row, tmp_path, monkeypatch):
    """sync.py sends only changed cards to Claude on a normal run, and every card with --review-all."""
    for mod in (build, sync):
        for name in ("SNAPSHOT_FILE", "EDITS_FILE", "ALIASES_FILE", "DECISIONS_FILE", "CARDS_FILE"):
            if hasattr(mod, name):
                monkeypatch.setattr(mod, name, tmp_path / getattr(build, name).name)
    monkeypatch.setattr(sync, "REVIEWS_FILE", tmp_path / "reviews.json")
    monkeypatch.setattr(sync, "REPORTS", tmp_path / "reports")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    sent = []

    class FakeReviewer:
        # Records which cards would be sent to Claude and "reviews" them with no findings.
        def __init__(self, *args, **kwargs):
            pass

        def review(self, cards):
            sent.append(sorted(c["hanzi"] for c in cards))
            return {c["id"]: {"hash": review.review_hash(c), "reviewed_at": "t", "hanzi": c["hanzi"],
                              "issues": [], "example": None, "example_error": ""} for c in cards}, ""

    monkeypatch.setattr(review, "Reviewer", FakeReviewer)
    tabs = _tabs(row)
    xlsx = tmp_path / "v.xlsx"
    xlsx.write_bytes(_xlsx(tabs))
    sync.main(["--xlsx", str(xlsx)])                     # first run: everything is unreviewed
    sync.main(["--xlsx", str(xlsx)])                     # nothing changed: nothing sent
    tabs["Places"][0]["notes"] = "line 2 goes downtown"  # change one card's content
    xlsx.write_bytes(_xlsx(tabs))
    sync.main(["--xlsx", str(xlsx)])
    sync.main(["--xlsx", str(xlsx), "--review-all"])
    assert sent == [["地铁", "杯子"], [], ["地铁"], ["地铁", "杯子"]]
